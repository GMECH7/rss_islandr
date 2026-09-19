import sys
from datetime import datetime
from tkinter import filedialog, messagebox

import pytest
import ttkbootstrap as tb
from openpyxl import load_workbook
from PIL import Image

from rss_islandr.core.datatypes import UICalcVariable, UIInpVariable
from rss_islandr.ui.io_btns import ExportExcelReportBtn

POLYGONS = (
    "[{'unique_id': 'a', 'type': 'polygon', 'name': 'P1', "
    "'coordinates': [[38.0, 54.5], [38.1, 54.6], [38.2, 54.5]], 'area_km2': 1.2, 'node_count': 3}]"
)


def _inp_var(value: str, text_val: str, excel_cell: str) -> UIInpVariable:
    return UIInpVariable(
        frame_tag="test_frame",
        tk_var=tb.StringVar(value=value),
        rel_pos=0,
        text_val=text_val,
        excel_cell=excel_cell,
    )


@pytest.fixture
def export(root, tmp_path, monkeypatch):
    """Run the excel export with prefilled variables. Returns a function that runs it and returns (file, errors)."""
    ui_inp_vars = {
        # Keys with the scenario id go only to that scenario's sheet, the others go to all three sheets
        "drop_on-on_IN_1_00": _inp_var("Large", "Extent", "C8"),
        "date_0_00": _inp_var("2025-09-19", "Assessment date", "C3"),
        "map_0_00": _inp_var("EPSG:4326", "CRS", "C4"),
        "map_0_01": _inp_var("54.5733", "Latitude", "C5"),
        "map_0_02": _inp_var("38.0048", "Longitude", "D5"),
        "site_0_00": _inp_var("", "Site name", "C2"),
    }
    ui_calc_vars = {"IN_on-on_frame": UICalcVariable(tb.StringVar(value="0.35"), "C9")}
    errors = []

    monkeypatch.setattr(messagebox, "showinfo", lambda *args, **kwargs: None)
    monkeypatch.setattr(messagebox, "showerror", lambda *args, **kwargs: errors.append(args))

    def run(image_files=(), polygons=POLYGONS, target=None):
        target = target or tmp_path / "report.xlsx"
        monkeypatch.setattr(messagebox, "askyesno", lambda *args, **kwargs: bool(image_files))
        monkeypatch.setattr(filedialog, "askopenfilenames", lambda **kwargs: [str(f) for f in image_files])
        monkeypatch.setattr(filedialog, "asksaveasfilename", lambda **kwargs: str(target))
        button = ExportExcelReportBtn(ui_inp_vars, ui_calc_vars, map_polygons_tb=tb.StringVar(value=polygons))
        button.on_btn_click()
        return target, errors

    return run


def test_values_are_written_with_the_right_types(export):
    """
    Steps:
    1. Prepare UI variables with text, an ISO date, numeric text, a risk value 0.35 and an empty text.
    2. Export the Excel report (no map images).

    Expected result:
    - In the first scenario sheet the text stays text, the date is a real date, the numeric text is a number and the
      risk 0.35 keeps the template's percentage format. The empty text leaves the cell empty.
    """
    target, errors = export()
    assert errors == []

    sheet = load_workbook(target).worksheets[0]
    assert sheet["C8"].value == "Large"
    assert sheet["C3"].value == datetime(2025, 9, 19)  # ISO date becomes a real date
    assert sheet["C5"].value == 54.5733  # numeric text becomes a number
    assert sheet["D5"].value == 38.0048
    assert sheet["C4"].value == "EPSG:4326"
    assert sheet["C9"].value == 0.35
    assert sheet["C9"].number_format == "0.00%"  # template format is kept
    assert sheet["C2"].value is None  # empty text is not written as text


def test_scenario_values_go_only_to_their_sheet(export):
    """
    Steps:
    1. Prepare a value of the on-site to on-site scenario and values that belong to no scenario.
    2. Export the Excel report.

    Expected result:
    - The scenario value is only in the first sheet. The other values are in all three scenario sheets.
    """
    target, _ = export()

    workbook = load_workbook(target)
    assert workbook.worksheets[0]["C8"].value == "Large"
    assert workbook.worksheets[1]["C8"].value is None
    assert workbook.worksheets[2]["C8"].value is None
    # Values without a scenario id are written to all three scenario sheets
    for sheet in workbook.worksheets[:3]:
        assert sheet["C5"].value == 54.5733


def test_polygons_are_written_to_the_coordinates_sheet(export):
    """
    Steps:
    1. Export the Excel report with one polygon (three nodes, area 1.2 km²).

    Expected result:
    - The coordinates sheet has the polygon name, the node numbers 1 to 3, the coordinates of a node and the area.
    """
    target, _ = export()

    sheet = load_workbook(target).worksheets[4]
    assert sheet["A2"].value == "P1"
    assert [sheet[f"B{row}"].value for row in (2, 3, 4)] == [1, 2, 3]
    assert (sheet["C3"].value, sheet["D3"].value) == (38.1, 54.6)
    assert sheet["E2"].value == 1.2


def test_map_images_are_added_below_each_other(export, tmp_path):
    """
    Steps:
    1. Create two PNG images of different sizes.
    2. Export the Excel report with both images.

    Expected result:
    - No error is shown. The maps sheet has two images, the first at the top and the second below it, and the aspect
      ratio of the first is kept.
    """
    image_files = []
    for i, size in enumerate([(800, 500), (600, 600)]):
        image_file = tmp_path / f"map{i}.png"
        Image.new("RGB", size, (70, 130, 180)).save(image_file)
        image_files.append(image_file)

    target, errors = export(image_files)
    assert errors == []

    images = load_workbook(target).worksheets[3]._images
    assert len(images) == 2
    first_row = images[0].anchor._from.row
    second_row = images[1].anchor._from.row
    assert first_row == 0
    assert second_row > first_row
    assert images[0].width / images[0].height == pytest.approx(800 / 500)


def test_error_is_reported_when_file_cannot_be_written(export, tmp_path):
    """
    Steps:
    1. Export the Excel report to a file in a folder that does not exist.

    Expected result:
    - No file is created and exactly one error message ('Error') is shown.
    """
    target, errors = export(target=tmp_path / "missing_folder" / "report.xlsx")

    assert not target.exists()
    assert len(errors) == 1
    assert errors[0][0] == "Error"


if __name__ == "__main__":
    # Enforce executing only this file verbose option and print statements in functions to be printed
    pytest.main([sys.argv[0], "-v", "-s"])
